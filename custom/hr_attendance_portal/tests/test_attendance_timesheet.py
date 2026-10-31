# -*- coding: utf-8 -*-

from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from odoo import fields
from odoo.exceptions import AccessError, UserError
from odoo.tests.common import TransactionCase, new_test_user, tagged

from .common import AttendancePortalMixin


@tagged("post_install", "-at_install")
class TestAttendanceTimesheet(AttendancePortalMixin, TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company.attendance_standard_hours = 8.0
        cls.company.attendance_overtime_cap_hours = 4.0
        cls.captain = new_test_user(
            cls.env,
            login="timesheet_captain",
            groups="hr_attendance_portal.group_hr_attendance_site_captain",
        )
        cls.captain.tz = "Asia/Ho_Chi_Minh"
        cls.captain.hr_authorized_department_ids = cls.department
        cls.support_department = cls.env["hr.department"].create({
            "name": "Support Team",
            "company_id": cls.company.id,
        })
        cls.other_department = cls.env["hr.department"].create({
            "name": "Other Lot",
            "company_id": cls.company.id,
        })
        cls.idle = cls.env["hr.employee"].create({
            "name": "An Idle",
            "emp_code": "KV-BCC-IDLE",
            "pin": "1357",
            "company_id": cls.company.id,
            "department_id": cls.department.id,
        })
        cls.employee.name = "Mai Official"
        cls.support = cls.env["hr.employee"].create({
            "name": "Binh Support",
            "emp_code": "KV-BCC-SUP",
            "pin": "1357",
            "company_id": cls.company.id,
            "department_id": cls.support_department.id,
        })

    def _shift(self, employee, start, hours, department):
        attendance = self.env["hr.attendance"].create({
            "employee_id": employee.id,
            "check_in": start,
            "in_department_id": department.id,
            "in_location_id": department.attendance_location_id.id,
        })
        attendance.with_context(attendance_original_punch=True).write({
            "check_out": start + timedelta(hours=hours),
        })
        return attendance

    def test_timesheet_splits_official_and_support(self):
        day = datetime(2026, 9, 1, 1, 0, 0)
        self._shift(self.employee, day, 8, self.department)
        denied = self._shift(self.employee, day + timedelta(days=1), 6, self.department)
        denied.action_deny_session("absent")
        overtime = self._shift(self.employee, day + timedelta(days=1, hours=6), 10, self.department)
        overtime.action_captain_overtime(2, "peak")
        self._shift(self.support, day, 4, self.department)
        self._shift(self.employee, day + timedelta(hours=8), 8, self.other_department)

        wizard = self.env["hr.attendance.timesheet.wizard"].with_user(self.captain).create({
            "department_id": self.department.id,
            "date_from": fields.Date.to_date("2026-09-01"),
            "date_to": fields.Date.to_date("2026-09-02"),
        })
        payload = wizard._prepare_timesheet()
        self.assertEqual([row["name"] for row in payload["official_day"]], ["An Idle", "Mai Official"])
        self.assertEqual(payload["official_night"], [])
        self.assertEqual([row["name"] for row in payload["support_day"]], ["Binh Support"])
        self.assertEqual(payload["support_night"], [])

        idle, official = payload["official_day"]
        self.assertEqual(idle["worked_days"], 0)
        self.assertEqual(idle["total"], 0)
        self.assertEqual(official["days"], [8, 10])
        self.assertEqual(official["regular"], 16)
        self.assertEqual(official["overtime"], 2)
        self.assertEqual(official["total"], 18)
        self.assertEqual(official["worked_days"], 2)
        self.assertEqual(official["shift"], "Ca Ngày")
        support = payload["support_day"][0]
        self.assertEqual(support["days"], [4, 0])
        self.assertEqual(support["regular"], 4)
        self.assertEqual(support["worked_days"], 1)
        self.assertEqual(support["support_site"], "Support Team")
        self.assertEqual(support["code"], "KV-BCC-SUP")

        workbook = load_workbook(BytesIO(wizard._render_xlsx()))
        sheet = workbook.active
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("T3", values)
        self.assertIn("T4", values)
        self.assertIn("Ca Ngày", values)
        self.assertIn("Ca Đêm", values)
        self.assertIn("Đến hỗ trợ Ca Ngày", values)
        self.assertIn("Đến hỗ trợ Ca Đêm", values)
        self.assertIn("Ghi chú Phụ Cấp", values)
        self.assertIn("Tổng giờ Công", values)
        weekday_above_first_day = False
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if getattr(value, "year", None) == 2026 and value.month == 9 and value.day == 1:
                    self.assertEqual(sheet.cell(cell.row - 1, cell.column).value, "T3")
                    weekday_above_first_day = True
        self.assertTrue(weekday_above_first_day)

    def test_export_rejects_an_inverted_or_long_range(self):
        wizard = self.env["hr.attendance.timesheet.wizard"].create({
            "department_id": self.department.id,
            "date_from": fields.Date.to_date("2026-09-02"),
            "date_to": fields.Date.to_date("2026-09-01"),
        })
        with self.assertRaises(UserError):
            wizard.action_export()
        wizard.date_to = fields.Date.to_date("2026-12-01")
        with self.assertRaises(UserError):
            wizard.action_export()

    def test_captain_cannot_export_another_lot(self):
        wizard = self.env["hr.attendance.timesheet.wizard"].with_user(self.captain).create({
            "department_id": self.other_department.id,
            "date_from": fields.Date.to_date("2026-09-01"),
            "date_to": fields.Date.to_date("2026-09-02"),
        })
        with self.assertRaises(AccessError):
            wizard.action_export()
